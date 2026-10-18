"""What a dataset's samples look like as an experiment: the factors that vary between samples, their levels with
counts, a cross-table of the two main factors, and what every sample shares. Read by code from the sample table a
repository offers: GEO sample characteristics, or a PRIDE SDRF or author sample sheet. No model is involved.
"""

import csv
import io
import json
import re
from collections import Counter

import httpx

from . import record, samples, sources

MAX_LEVELS = 12
NOT_STATED = re.compile(r"^(not available|not applicable|n/?a|na|none|unknown|-|)$", re.I)
# Columns that name files, runs or individual samples, or describe the measurement rather than the samples.
NOT_A_FACTOR = re.compile(
    r"file|uri|raw|source name|assay name|accession|^title|replicate|fraction|^comment\[|technology type|"
    r"sdrf|label$|instrument|modification|cleavage|dissociation|\bid\b|identifier|^name$|description",
    re.I,
)
CONDITION = re.compile(r"disease|diagnos|condition|treatment|genotype|status|group|phenotype|agent|stimul|infection|time|duration", re.I)
SHEET = re.compile(r"\.sdrf\.tsv$|sample|design|metadata|annotation|condition|group|experiment", re.I)
SHEET_TYPES = (".tsv", ".csv", ".txt", ".xlsx")
MAX_SHEET_BYTES = 2_000_000


def _clean(name: str) -> str:
    """`characteristics[cell type]` and `factor value[cell type]` both become `cell type`."""
    return re.sub(r"^(characteristics|factor value)\[(.*)\]$", r"\2", name.strip(), flags=re.I).strip()


def _partition(values: list) -> tuple[int, ...]:
    """Which samples share a value, independent of how the values are written."""
    first: dict = {}
    return tuple(first.setdefault(v, len(first)) for v in values)


def from_rows(rows: list[dict[str, str]], source: str) -> dict | None:
    """Factors are columns that split the samples into 2 to 12 groups; columns with one value are shared context."""
    if len(rows) < 2:
        return None
    # A sample without a value for a column still takes a place in it, so every column lines up with the rows.
    column_of: dict[str, str] = {}
    for name in dict.fromkeys(k for row in rows for k in row if k and not NOT_A_FACTOR.search(k)):
        column_of.setdefault(_clean(name), name)
    columns = {name: [str(row.get(column) or "").strip() for row in rows] for name, column in column_of.items()}
    factors, shared = [], {}
    for name, values in columns.items():
        stated = [v for v in values if not NOT_STATED.match(v)]
        levels = Counter(stated)
        if len(levels) == 1 and len(stated) == len(rows):
            shared[name] = stated[0]
        elif 2 <= len(levels) <= MAX_LEVELS and len(levels) < len(rows):
            factors.append({"name": name, "column": column_of[name], "levels": dict(levels.most_common()),
                            **({"unstated": len(rows) - len(stated)} if len(stated) < len(rows) else {})})
    # Submitters often record one grouping under several names, or join two factors into a combined column
    # ("IPF.Base"); either adds nothing, so duplicates become aliases and combinations are dropped.
    unique: dict[tuple, dict] = {}
    for f in factors:
        key = _partition(columns[f["name"]])
        if key in unique:
            unique[key].setdefault("aliases", []).append(f["name"])
        else:
            unique[key] = f
    names = [f["name"] for f in unique.values()]
    # A combination is finer than both of its parts; a factor nested in another reproduces itself and is kept.
    combined = {key for a in names for b in names if a < b
                if (key := _partition(list(zip(columns[a], columns[b])))) not in (_partition(columns[a]), _partition(columns[b]))}
    factors = [f for key, f in unique.items() if key not in combined]
    if not factors and not shared:
        return None
    # Condition-like factors lead, then those with the fewest levels; the first two form the cross-table.
    factors.sort(key=lambda f: (not CONDITION.search(f["name"]), len(f["levels"])))
    design = {"source": source, "samples": len(rows), "factors": factors, "shared": shared}
    if len(factors) >= 2:
        a, b = factors[0]["name"], factors[1]["name"]
        pairs = Counter((columns[a][i], columns[b][i]) for i in range(len(rows)))
        design["crossed"] = {"rows": a, "columns": b, "counts": [[x, y, n] for (x, y), n in pairs.items()]}
    return design


def _geo(accession: str) -> dict | None:
    records = samples.geo_records(accession)
    rows = [{k: v for k, v in r["fields"].items() if k != "title pattern"} for r in records]
    return from_rows(rows, "GEO sample characteristics")


def _sheet_rows(url: str) -> list[dict[str, str]]:
    path = samples.CACHE / "sheets" / re.sub(r"[^\w.-]", "_", url.rsplit("/", 2)[-2] + "_" + url.rsplit("/", 1)[-1])
    if not path.exists():
        path.parent.mkdir(parents=True, exist_ok=True)
        path.write_bytes(httpx.get(url, timeout=60, follow_redirects=True).raise_for_status().content)
    if path.suffix == ".xlsx":
        from openpyxl import load_workbook

        sheet = load_workbook(path, read_only=True, data_only=True).worksheets[0]
        table = [["" if c is None else str(c) for c in row] for row in sheet.iter_rows(values_only=True)]
    else:
        text = path.read_text(errors="replace")
        table = list(csv.reader(io.StringIO(text), delimiter="\t" if "\t" in text.split("\n", 1)[0] else ","))
    table = [r for r in table if any(c.strip() for c in r)]
    if len(table) < 3:
        return []
    header = [h.strip() for h in table[0]]
    # Proteomics tables list one row per raw file; a sample measured in several fractions repeats.
    rows = [dict(zip(header, r)) for r in table[1:]]
    key = next((h for h in header if h.lower() == "source name"), None)
    return list({r[key]: r for r in rows}.values()) if key else rows


def _pride(accession: str) -> dict | None:
    sources.pride_files(accession)
    files = json.loads((sources.PRIDE_FILES_CACHE / f"{accession}.json").read_text())
    candidates = []
    for f in files:
        name, category = f["fileName"], f["fileCategory"]["value"]
        url = next((loc["value"] for loc in f.get("publicFileLocations", []) if loc["value"].startswith("ftp://")), None)
        if url and f.get("fileSizeBytes", 0) <= MAX_SHEET_BYTES and name.lower().endswith(SHEET_TYPES) and (category == "EXPERIMENTAL DESIGN" or SHEET.search(name)):
            candidates.append((name, url.replace("ftp://", "https://", 1)))
    designs = []
    for name, url in candidates:
        try:
            if design := from_rows(_sheet_rows(url), name):
                designs.append(design)
        except (httpx.HTTPError, ValueError, KeyError, OSError):
            continue
    # Community SDRFs are often minimal; the sheet that separates samples into the most factors is the informative one.
    return max(designs, key=lambda d: (len(d["factors"]), len(d["shared"])), default=None)


# Columns of a record's sample table that identify runs or measure output; neither describes the samples.
NOT_SAMPLE = {"Run", "Sample", "Reads", "Bases", "Experiment"}


def _record(accession: str, source: str) -> dict | None:
    """ENA run tables and ArrayExpress characteristics tables, from the dataset's record. ArrayExpress lists each
    combination of characteristics once with its number of samples, so rows are repeated by that number."""
    row = record._row(accession)
    table = row and record.raw(accession, row).samples
    if not table:
        return None
    rows = [dict(zip(table["columns"], r)) for r in table["rows"]]
    rows = [{k: v for k, v in r.items() if k not in NOT_SAMPLE and k != "No. of Samples"}
            for r in rows for _ in range(int(r.get("No. of Samples") or 1))]
    label = {"MetabolomicsWorkbench": "Metabolomics Workbench sample factors", "PRIDE": "sample names in the authors' protein table"}.get(source, f"{source} sample table")
    return from_rows(rows, label)


def design(accession: str, source: str) -> dict | None:
    if source == "GEO":
        return _geo(accession)
    if source in ("PRIDE", "iProX", "jPOST", "PanoramaPublic") and accession.startswith("PXD"):
        return _pride(accession) or _record(accession, "PRIDE")
    if source in ("ENA", "ArrayExpress", "MetabolomicsWorkbench", "MetaboLights"):
        return _record(accession, source)
    return None


def people(accession: str, comparison: dict | None) -> dict | None:
    """Distinct patients or donors in the sample table, in total and per compared group, when a column names them."""
    rec = record.record(accession)
    if not rec or not rec["samples"] or not rec["samples"]["rows"]:
        return None
    cols, rows = rec["samples"]["columns"], rec["samples"]["rows"]
    i = next((n for n, c in enumerate(cols) if samples.DONOR.match(c.strip())), None)
    named = [r for r in rows if i is not None and r[i]]
    if i is None or not 2 <= len({r[i] for r in named}) < len(rows):
        return None
    out = {"field": cols[i], "total": len({r[i] for r in named}), "samples": len(rows)}
    if comparison and comparison["factor"] in cols:
        f = cols.index(comparison["factor"])
        out |= {g: len({r[i] for r in named if r[f] in comparison[g]}) for g in ("case", "control")}
    return out


if __name__ == "__main__":
    rows = [{"Raw datafile name": f"x_{i}.raw", "Donor ": d, "Condition": c, "Duration, hours": h}
            for i, (d, c, h) in enumerate([(d, c, h) for d in ("F53", "M65", "F40") for c in ("Normoxia", "Hypoxia") for h in ("24", "48")])]
    built = from_rows(rows, "sheet")
    assert [f["name"] for f in built["factors"]] == ["Condition", "Duration, hours", "Donor"], built
    assert built["crossed"]["rows"] == "Condition" and [3, 3] == sorted(n for x, y, n in built["crossed"]["counts"] if x == "Hypoxia"), built
    sdrf = [{"source name": f"s{i}", "characteristics[organism]": "Homo sapiens", "characteristics[disease]": "normal",
             "comment[data file]": f"f{i}.raw", "characteristics[sex]": "not available"} for i in range(4)]
    minimal = from_rows(sdrf, "sdrf")
    assert minimal["factors"] == [] and minimal["shared"] == {"organism": "Homo sapiens", "disease": "normal"}, minimal
    twice = from_rows([r | {"Treatment": r["Condition"], "Group": r["Condition"] + "." + r["Duration, hours"]} for r in rows], "sheet")
    assert [f["name"] for f in twice["factors"]] == ["Condition", "Duration, hours", "Donor"] and twice["factors"][0]["aliases"] == ["Treatment"], twice
    uneven = from_rows([{"Condition": "a", "Dose": "1"}, {"Condition": "b"}, {"Condition": "a", "Dose": "2"}, {"Condition": "b", "Dose": "1"}], "sheet")
    assert uneven["samples"] == 4 and all(sum(f["levels"].values()) <= 4 for f in uneven["factors"]), uneven
    for accession, factors in (("GSE179525", ["genotype", "age", "library strategy"]), ("GSE149204", ["library strategy", "cell type", "molecule"])):
        assert [f["name"] for f in design(accession, "GEO")["factors"]] == factors, accession
    print("ok")
